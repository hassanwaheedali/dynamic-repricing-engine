"""Unit and integration tests for 5% Personal Commission + Winning Products Finder."""

from __future__ import annotations

import asyncio
from pathlib import Path
from unittest.mock import AsyncMock, patch

import openpyxl
import pandas as pd

from index import main, parse_args, prompt_run_mode
from pricing_engine.calculator import (
    calculate_effective_min_profit,
    calculate_effective_target_profit,
    calculate_floors,
    calculate_personal_commission,
    split_real_profit,
)
from pricing_engine.exporter import export_excel_workbook
from pricing_engine.models import PlatformPricingResult, ScrapedOffer, ScrapeResult
from pricing_engine.pipeline import (
    determine_risk_alert,
    parse_product_input,
    process_product_with_offers,
    run_pipeline,
)


def test_reusable_calculator_functions() -> None:
    # 1. calculate_effective_min_profit: 500 / 0.95 = 526.315789...
    eff_min = calculate_effective_min_profit(500.0, 5.0)
    assert round(eff_min, 2) == 526.32
    assert calculate_effective_min_profit(500.0, 0.0) == 500.0

    # 2. calculate_effective_target_profit: ceil(500 / 0.95) = 527
    assert calculate_effective_target_profit(500.0, 5.0) == 527
    assert calculate_effective_target_profit(1000.0, 5.0) == 1000
    assert calculate_effective_target_profit(None, 5.0) == 527
    assert calculate_effective_target_profit(500.0, 0.0) == 500

    # 3. split_real_profit: splits into (user_commission, company_profit)
    assert split_real_profit(1000.0, 5.0) == (50, 950)
    assert split_real_profit(527.0, 5.0) == (26, 501)
    assert split_real_profit(None, 5.0) == (0, None)
    assert split_real_profit(-500.0, 5.0) == (0, -500)
    assert split_real_profit(1000.0, 0.0) == (0, 1000)


def test_calculate_personal_commission_on_real_profit_only() -> None:
    # 5% of Real Net Profit (1,000) = 50, 5% of 4,000 = 200
    assert calculate_personal_commission(1000.0, 5.0) == 50
    assert calculate_personal_commission(4000.0, 5.0) == 200
    assert calculate_personal_commission(1000.0, 0.0) == 0
    assert calculate_personal_commission(0.0, 5.0) == 0


def test_calculate_floors_with_5pct_personal_commission() -> None:
    # Cost Base = 102,000. Min Company Profit = 500.
    # Total Min Profit at floor = 500 / 0.95 = 526.315...
    # Alfa Floor (10%) = (102,000 + 526.315) / 0.90 = 113,918.13
    # Digi 8% Floor = (102,000 + 526.315) / 0.92 = 111,441.65
    floors = calculate_floors(
        cost_base=102000.0,
        min_profit=500.0,
        net_price=100000.0,
        personal_comm_pct=5.0,
    )
    assert round(floors.floor_alfa) == 113918
    assert round(floors.floor_digi_10) == 113918
    assert round(floors.floor_digi_8) == 111442
    assert floors.personal_comm == 26


def test_process_product_5pct_win_and_columns() -> None:
    # Net Price = 100,000, Shipping = 2,000 -> Cost Base = 102,000
    # Floor: Alfa 113,918 | Digi 8% 111,442
    product = parse_product_input(
        0,
        pd.Series(
            {
                "S.NO": 1,
                "Brand": "Dawlance",
                "Name": "Dawlance Auto Washer DWT-1165 PL",
                "Shipping Estimate": 2000,
                "Net Price": 100000,
            }
        ),
    )
    alfa_offers = [
        ScrapedOffer("Seller A", 122000.0, False, "Dawlance DWT-1165 PL"),
        ScrapedOffer("Seller B", 124500.0, False, "Dawlance DWT-1165 PL"),
    ]
    digi_offers = [
        ScrapedOffer("Digi Seller 1", 118500.0, False, "Dawlance DWT-1165 PL"),
        ScrapedOffer("Digi Seller 2", 121000.0, False, "Dawlance DWT-1165 PL"),
        ScrapedOffer("Digi Seller 3", 123000.0, False, "Dawlance DWT-1165 PL"),
    ]

    rec = process_product_with_offers(product, alfa_offers, digi_offers, personal_comm_pct=5.0)

    # Risk Alerts must still be the final two keys
    keys = list(rec.keys())
    assert keys[-2:] == ["Alfa Risk Alert", "Digi Risk Alert"]

    # Verify 5% Personal Commission and Company Profit breakdown
    assert rec["Alfa Active Sellers"] == 2
    assert rec["Digi Active Sellers"] == 3
    assert rec["Min Profit Target (5% + 500)"] == 527

    alfa_p = rec["Alfa Real Profit"]
    assert alfa_p is not None and alfa_p >= 527
    assert rec["Alfa Company Profit"] == alfa_p - round(alfa_p * 0.05)
    assert rec["Alfa Company Profit"] >= 500
    assert rec["Alfa 5% Goal"] == "WIN + 5% ACHIEVED"

    digi_p = rec["Digi Real Profit"]
    assert digi_p is not None and digi_p >= 527
    assert rec["Digi Company Profit"] == digi_p - round(digi_p * 0.05)
    assert rec["Digi Company Profit"] >= 500
    assert rec["Digi 5% Goal"] == "WIN + 5% ACHIEVED"
    assert rec["My 5% Commission"] == max(round(alfa_p * 0.05), round(digi_p * 0.05))


def test_process_product_5pct_rejects_competitor_below_5pct_floor() -> None:
    # Competitor is at 108,000, which is below our floor (113,918 on Alfa, 111,442 on Digi).
    # Must NOT undercut at the expense of our floor!
    product = parse_product_input(
        0,
        pd.Series(
            {
                "S.NO": 1,
                "Brand": "Dawlance",
                "Name": "Dawlance Auto Washer DWT-1165 PL",
                "Shipping Estimate": 2000,
                "Net Price": 100000,
            }
        ),
    )
    cheap_offers = [ScrapedOffer("Cheap Seller", 108000.0, False, "Dawlance DWT-1165 PL")]

    rec = process_product_with_offers(product, cheap_offers, cheap_offers, personal_comm_pct=5.0)

    assert rec["Alfa Status"] == "FALLBACK"
    assert rec["Digi Status"] == "FALLBACK"
    assert rec["Alfa 5% Goal"] == "FALLBACK (UNCOMPETITIVE)"
    assert rec["Digi 5% Goal"] == "FALLBACK (UNCOMPETITIVE)"
    assert rec["Alfa Real Profit"] >= 500
    assert rec["Digi Real Profit"] >= 500


def test_export_winning_5pct_results_filters_and_sorts_by_demand(tmp_path: Path) -> None:
    input_csv = tmp_path / "catalog.csv"
    out_dir = tmp_path / "outputs"

    df_in = pd.DataFrame(
        [
            {
                "S.NO": 1,
                "Brand": "Gree",
                "Name": "Gree Split AC GS-18AITH21W",
                "Length": 90,
                "Width": 30,
                "Height": 30,
                "Net Price": 100000,
            },
            {
                "S.NO": 2,
                "Brand": "Haier",
                "Name": "Haier Split AC HSU-18HFN",
                "Length": 90,
                "Width": 30,
                "Height": 30,
                "Net Price": 120000,
            },
            {
                "S.NO": 3,
                "Brand": "Dawlance",
                "Name": "Dawlance Fridge 9178LF",
                "Length": 60,
                "Width": 60,
                "Height": 150,
                "Net Price": 80000,
            },
        ]
    )
    df_in.to_csv(input_csv, index=False)

    async def _mock_scrape_alfa(session, patterns, target_model, **kwargs):
        if "18AITH21W" in target_model:
            # 1 seller, viable above 5% floor
            return ScrapeResult(
                offers=[ScrapedOffer("V1", 125000.0, False, "Gree GS-18AITH21W")]
            )
        if "18HFN" in target_model:
            # 4 sellers (Higher Demand!), viable above 5% floor
            return ScrapeResult(
                offers=[
                    ScrapedOffer("V1", 148000.0, False, "Haier HSU-18HFN"),
                    ScrapedOffer("V2", 149000.0, False, "Haier HSU-18HFN"),
                    ScrapedOffer("V3", 150000.0, False, "Haier HSU-18HFN"),
                    ScrapedOffer("V4", 152000.0, False, "Haier HSU-18HFN"),
                ]
            )
        # Item 3: Competitor dumping below floor -> NOT a winner
        return ScrapeResult(
            offers=[ScrapedOffer("DumpingVendor", 82000.0, False, "Dawlance 9178LF")]
        )

    async def _mock_scrape_digi(session, patterns, target_model, **kwargs):
        return ScrapeResult(offers=[])

    with (
        patch("pricing_engine.pipeline.scrape_alfamall", side_effect=_mock_scrape_alfa),
        patch("pricing_engine.pipeline.scrape_digimall", side_effect=_mock_scrape_digi),
    ):
        out_df = asyncio.run(
            run_pipeline(
                input_path=input_csv,
                export_dir=out_dir,
                concurrency=2,
                verbose=True,
                personal_comm_pct=5.0,
            )
        )

    assert len(out_df) == 3
    winners_csv = out_dir / "catalog_winning_5pct_latest.csv"
    winners_xlsx = out_dir / "catalog_winning_5pct_latest.xlsx"
    assert winners_csv.is_file()
    assert winners_xlsx.is_file()

    w_df = pd.read_csv(winners_csv)
    # Only items 1 and 2 won with 5% + 500; item 2 has 4 sellers vs item 1's 1 seller so item 2 is ranked #1!
    assert len(w_df) == 2
    assert list(w_df["S.NO"]) == [2, 1]
    assert w_df["My 5% Commission"].iloc[0] > 0
    assert w_df["My 5% Commission"].iloc[1] > 0

    # Verify Excel formulas in winners workbook
    wb = openpyxl.load_workbook(winners_xlsx)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    col_map = {h: openpyxl.utils.get_column_letter(i + 1) for i, h in enumerate(headers)}
    my_comm_formula = ws[f"{col_map['My 5% Commission']}2"].value
    assert isinstance(my_comm_formula, str) and "*0.05" in my_comm_formula


def test_cli_personal_comm_flags_and_interactive_mode(tmp_path: Path) -> None:
    with patch("sys.argv", ["index.py", "--5pct"]):
        args = parse_args()
        assert args.five_pct is True

    with patch("sys.argv", ["index.py", "-m", "5"]):
        args = parse_args()
        assert args.personal_comm == 5.0

    with patch("builtins.input", return_value="3"):
        assert prompt_run_mode() == "reprice_5pct"

    test_csv = tmp_path / "test.csv"
    test_csv.write_text("Brand,Name,Net Price\nGree,GS-18AITH21W,100000\n", encoding="utf-8")
    with (
        patch("sys.argv", ["index.py", "--5pct", "-i", str(test_csv)]),
        patch("index.run_pipeline", new_callable=AsyncMock) as mock_pipe,
    ):
        assert asyncio.run(main()) == 0
        mock_pipe.assert_called_once()
        assert mock_pipe.call_args.kwargs.get("personal_comm_pct") == 5.0


def test_adversarial_edge_cases(tmp_path: Path) -> None:
    # 1. Explicit personal_comm field on CostFloors: custom min_profit without personal_comm_pct
    floors_no_comm = calculate_floors(cost_base=100000.0, min_profit=1200.0, personal_comm_pct=0.0)
    assert floors_no_comm.personal_comm == 0

    floors_with_comm = calculate_floors(
        cost_base=100000.0, min_profit=500.0, net_price=100000.0, personal_comm_pct=3.5
    )
    # 500 / (1 - 0.035) - 500 = 518.13 - 500 = 18
    assert floors_with_comm.personal_comm == 18

    # 2. THIN_BUFFER triggers when company profit (real_profit - personal_comm) < 700
    res = PlatformPricingResult(
        min_price=120000.0,
        max_price=125000.0,
        final_price=119899,
        real_profit=5600,  # 5000 personal comm + 600 company profit (< 700 buffer)
        status="Competitive Win",
        reason="Undercut lowest competitor",
    )
    assert determine_risk_alert(res, personal_comm=5000) == "THIN_BUFFER"
    assert determine_risk_alert(res, personal_comm=4800) == "NONE"

    # 3. Dynamic personal_comm_pct in Excel formula (e.g. 3.5% -> *0.035)
    xlsx_path = tmp_path / "custom_comm.xlsx"
    df = pd.DataFrame(
        [
            {
                "S.NO": 1,
                "Brand": "Gree",
                "Name": "GS-18AITH21W",
                "Net Price": 100000,
                "My 5% Commission": 3500,
                "Min Profit Target (5% + 500)": 4000,
            }
        ]
    )
    export_excel_workbook(df, xlsx_path, personal_comm_pct=3.5)
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    col_map = {h: openpyxl.utils.get_column_letter(i + 1) for i, h in enumerate(headers)}
    formula = ws[f"{col_map['My 5% Commission']}2"].value
    assert isinstance(formula, str) and "*0.035" in formula
