"""Tests for dynamic Excel formula generation and dual CSV/XLSX export."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

from pricing_engine.exporter import (
    ExportManifest,
    export_results,
)


def test_export_results_dual_artifacts_with_formulas(tmp_path: Path) -> None:
    out_dir = tmp_path / "outputs"
    df = pd.DataFrame(
        {
            "S.NO": [1, 2],
            "Brand": ["Haier", "Gree"],
            "Name": ["Haier Washer 15KG", "Gree Split AC 1.5T"],
            "Net Price": [100000, 140000],
            "Length": [60.0, 100.0],
            "Width": [60.0, 40.0],
            "Height": [85.0, 90.0],
            "Profit": [1000, 1500],
            "Shipping Estimate": [4284, 5040],
            "Commision": ["10.00%", "10.00%"],
            "Digi Effective Commision": ["8.00%", "10.00%"],
            "Total Without Profit": [113352, 161156],
            "Calculation": [117000, 162800],
            "Minimum Price Alfa": [116427, 161711],
            "Maximum Price Alfa": [140000, 190000],
            "Minimum Price Digi": [113896, 161711],
            "Maximum Price Digi": [140000, 190000],
            "Final Price Alfa": [122899, 168999],
            "Alfa Status": ["WINNER", "WINNER"],
            "Final Price Digi": [121899, 167899],
            "Digi Status": ["WINNER", "WINNER"],
            "Alfa Real Profit": [6325, 7059],
            "Digi Real Profit": [7863, 6069],
            "Alfa Reason": ["Undercut lowest competitor", "Undercut lowest competitor"],
            "Digi Reason": ["Undercut lowest competitor", "Undercut lowest competitor"],
            "Alfa Risk Alert": ["NONE", "NONE"],
            "Digi Risk Alert": ["NONE", "NONE"],
        }
    )

    fixed_dt = datetime(2026, 10, 2, 10, 30, 0)
    manifest = export_results(
        df=df,
        source_filename="catalog_input.csv",
        output_dir=out_dir,
        timestamp=fixed_dt,
        render_panel=False,
    )

    assert isinstance(manifest, ExportManifest)
    assert manifest.total_rows == 2

    # 1. Verify both CSV artifacts exist
    assert manifest.timestamped_path.is_file()
    assert manifest.latest_path.is_file()
    assert manifest.timestamped_path.name == "catalog_input_repriced_20261002_103000.csv"
    assert manifest.latest_path.name == "catalog_input_repriced_latest.csv"

    # 2. Verify both XLSX artifacts exist
    assert manifest.timestamped_xlsx_path is not None
    assert manifest.latest_xlsx_path is not None
    assert manifest.timestamped_xlsx_path.is_file()
    assert manifest.latest_xlsx_path.is_file()
    assert manifest.timestamped_xlsx_path.name == "catalog_input_repriced_20261002_103000.xlsx"
    assert manifest.latest_xlsx_path.name == "catalog_input_repriced_latest.xlsx"

    # 3. Load Excel workbook and verify dynamic formulas
    wb = openpyxl.load_workbook(manifest.latest_xlsx_path)
    ws = wb.active
    assert ws.title == "Repriced Catalog"

    # Verify freeze pane
    assert ws.freeze_panes == "A2"

    # Read row 2 (first product) formulas
    headers = [cell.value for cell in ws[1]]
    col_map = {h: openpyxl.utils.get_column_letter(i + 1) for i, h in enumerate(headers)}

    # Shipping Estimate formula check
    ship_val = ws[f"{col_map['Shipping Estimate']}2"].value
    assert isinstance(ship_val, str) and ship_val.startswith("=ROUNDUP")
    assert f"{col_map['Length']}2" in ship_val
    assert f"{col_map['Width']}2" in ship_val
    assert f"{col_map['Height']}2" in ship_val

    # Total Without Profit formula check
    twop_val = ws[f"{col_map['Total Without Profit']}2"].value
    assert isinstance(twop_val, str) and twop_val.startswith("=ROUND")
    assert f"{col_map['Net Price']}2" in twop_val
    assert f"{col_map['Shipping Estimate']}2" in twop_val

    # Calculation formula check
    calc_val = ws[f"{col_map['Calculation']}2"].value
    assert isinstance(calc_val, str) and calc_val.startswith("=ROUND")
    assert f"{col_map['Profit']}2" in calc_val

    # Minimum Price Alfa & Digi check (real scraped market numbers, NOT formulas)
    min_alfa_val = ws[f"{col_map['Minimum Price Alfa']}2"].value
    min_digi_val = ws[f"{col_map['Minimum Price Digi']}2"].value
    assert min_alfa_val == 116427
    assert min_digi_val == 113896

    # Real Profit formulas check
    prof_alfa_val = ws[f"{col_map['Alfa Real Profit']}2"].value
    prof_digi_val = ws[f"{col_map['Digi Real Profit']}2"].value
    assert isinstance(prof_alfa_val, str) and "0.90" in prof_alfa_val
    assert isinstance(prof_digi_val, str) and f"{col_map['Digi Effective Commision']}2" in prof_digi_val

    # 4. Verify CSV content remains clean integer format for portal upload
    csv_df = pd.read_csv(manifest.latest_path, dtype=str)
    assert csv_df["Final Price Alfa"].iloc[0] == "122899"
    assert csv_df["Total Without Profit"].iloc[0] == "113352"
    assert csv_df["Calculation"].iloc[0] == "117000"
    assert not csv_df["Total Without Profit"].iloc[0].startswith("=")
