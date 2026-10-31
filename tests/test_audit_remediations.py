"""Comprehensive regression tests for all independent engineering audit findings.

Covers Critical (AUD-CRIT-01..06), High (AUD-HIGH-01..10), Medium (AUD-MED-01..07),
Low (AUD-LOW-01..04), Newly Discovered (AUD-NEW-01..03), and cross-brand Pakistani
appliance isolation across all 13 target brands.
"""

from __future__ import annotations

import asyncio
import math
from pathlib import Path
from typing import Any
from unittest.mock import AsyncMock, MagicMock, patch

import openpyxl
import pandas as pd
import pytest

import config
from pricing_engine.calculator import (
    calculate_cost_base,
    calculate_effective_target_profit,
    calculate_floors,
    calculate_shipping_estimate,
    calculate_volumetric_weight,
)
from pricing_engine.checkpoint import (
    compute_row_fingerprint,
    detect_dirty_or_pending_rows,
)
from pricing_engine.db import SupabaseClient
from pricing_engine.evaluator import (
    evaluate_alfa_pricing,
    evaluate_digi_pricing,
    format_fallback_descriptions,
)
from pricing_engine.exporter import (
    export_results,
    export_winning_5pct_results,
)
from pricing_engine.learner import AutonomousLearner
from pricing_engine.matcher import (
    clean_marketplace_noise,
    expand_cartesian_branches,
    extract_product_slots,
    extract_search_query_and_model,
    get_capacity_map,
    get_capacity_tonnage,
    referee_match,
)
from pricing_engine.models import (
    AmbiguousVariantConflict,
    CheckpointState,
    ProductInput,
    ScrapedOffer,
    ScrapeResult,
    is_non_competitor_vendor,
)
from pricing_engine.pipeline import (
    _format_summary_price,
    default_learner,
    parse_product_input,
    process_item_async,
    process_product_with_offers,
    run_pipeline,
)
from pricing_engine.range_pipeline import (
    load_range_products,
    run_range_pipeline,
)
from pricing_engine.registry import InMemoryRegistry
from pricing_engine.scraper import (
    _query_cache,
    _raw_items_cache,
    _raw_items_meta_cache,
    clear_scraper_cache,
    deduplicate_offers,
    parse_price_value,
    scrape_alfamall,
    scrape_digimall,
)


class _DummyResponse:
    """Typed mock HTTP response for curl_cffi AsyncSession."""

    def __init__(self, status_code: int, payload: Any) -> None:
        self.status_code = status_code
        self._payload = payload

    def json(self) -> Any:
        return self._payload


def _assert_match(
    target_title: str,
    candidate_title: str,
    brand: str,
    expected: bool,
) -> Any:
    t_slots = extract_product_slots(target_title, brand=brand)
    dec = referee_match(t_slots, candidate_title, brand=brand)
    assert dec.is_match is expected, (
        f"Expected is_match={expected} for target={target_title!r} vs "
        f"candidate={candidate_title!r} (brand={brand!r}), got {dec}"
    )
    return dec


# ==============================================================================
# CRITICAL FINDINGS (AUD-CRIT-01 .. AUD-CRIT-06)
# ==============================================================================


def test_aud_crit_01_unprefixed_glued_variants() -> None:
    """AUD-CRIT-01: Unprefixed glued 3/4-letter variant suffixes must split and collide."""
    _assert_match("Haier HRF-216 EBD Refrigerator", "Haier 216EBS Refrigerator", "Haier", False)
    _assert_match("Haier HRF-306 EPR Refrigerator", "Haier 306EPB Refrigerator", "Haier", False)
    _assert_match("Haier HRF-368 IPRA Refrigerator", "Haier 368IPGA Refrigerator", "Haier", False)

    # Identical variant suffix matches cleanly whether prefixed or unprefixed
    dec = _assert_match(
        "Haier HRF-216 EBS Refrigerator", "Haier 216EBS Refrigerator", "Haier", True
    )
    assert dec.is_base_match is False

    # Metric units glued to stems must NOT be misclassified as variant tags
    cft_slots = extract_product_slots("Dawlance 9178 14CFT Refrigerator", brand="Dawlance")
    assert "-CFT" not in cft_slots.guarded_facets.capabilities
    assert "14.0_CFT" not in cft_slots.guarded_facets.capabilities


def test_aud_crit_02_alphanumeric_compound_anchors() -> None:
    """AUD-CRIT-02: Single-number anchors must reject alphanumeric sub-model compound stems."""
    _assert_match("Haier HWM 120 Washing Machine", "Haier HWM 120-B699S8 Top Load", "Haier", False)
    _assert_match("Haier HWM 120 Washing Machine", "Haier HWM 120-M826 Automatic", "Haier", False)
    _assert_match(
        "Dawlance DW 260 Washing Machine", "Dawlance DW 260-X500 Washer", "Dawlance", False
    )
    _assert_match("Haier HWM 80 Washing Machine", "Haier HWM 80-1708 Washer", "Haier", False)


def test_aud_crit_03_preserve_hardware_specs_after_hyphen_and_pipe() -> None:
    """AUD-CRIT-03: Hardware specs after ' - ' or ' | ' must survive while seller noise is stripped."""
    # Cool Only vs Heat & Cool after ' - '
    _assert_match(
        "Gree GS-18PITH11W Heat & Cool Split AC",
        "Gree GS-18PITH11W Split AC - Cool Only",
        "Gree",
        False,
    )
    # Twin Tub after ' - ' vs Front Load
    _assert_match(
        "Haier HWM-100 Front Load Washing Machine",
        "Haier HWM-100 Washing Machine - Twin Tub",
        "Haier",
        False,
    )
    # Variant suffix '- EBS' vs 'EBD'
    _assert_match(
        "Haier HRF-276 EBD Refrigerator",
        "Haier HRF-276 Refrigerator - EBS",
        "Haier",
        False,
    )
    # Door finish 'Metal Door' vs 'Glass Door'
    _assert_match(
        "Dawlance 9178 LF Glass Door Refrigerator",
        "Dawlance 9178 LF Refrigerator - Metal Door",
        "Dawlance",
        False,
    )
    # Multi-pipe: middle spec preserved, trailing seller clause stripped
    cleaned = clean_marketplace_noise(
        "Haier HSU-18HFP | T3 Inverter | Official Store Free Delivery"
    )
    assert "T3 Inverter" in cleaned
    assert "Official Store" not in cleaned
    slots = extract_product_slots(
        "Haier HSU-18HFP | T3 Inverter | Official Store Free Delivery", brand="Haier"
    )
    assert "T3" in slots.guarded_facets.capabilities
    assert slots.guarded_facets.is_inverter is True


def test_aud_crit_04_explicit_sku_prefix_collisions() -> None:
    """AUD-CRIT-04: Explicit distinct SKU prefixes sharing brand[0] must never cross-match."""
    _assert_match("Haier HRF-240 Refrigerator", "Haier HWM-240 Washing Machine", "Haier", False)
    _assert_match("Dawlance DW-200 Microwave", "Dawlance DMB-200 Mini Bar", "Dawlance", False)
    _assert_match("Dawlance DWT-260 Washer", "Dawlance DF-260 Deep Freezer", "Dawlance", False)
    _assert_match("Super Asia SA-240 Washer", "Super Asia SD-240 Spinner", "Super Asia", False)
    _assert_match("Westpoint WF-500 Juicer", "Westpoint WM-500 Mixer", "Westpoint", False)

    # Asymmetric unprefixed target vs valid brand-prefixed candidate still matches
    _assert_match("Dawlance 9178WB Refrigerator", "Dawlance REF DW-9178WB", "Dawlance", True)


def test_aud_crit_05_empty_anchor_and_form_factor_guards() -> None:
    """AUD-CRIT-05: Model-less titles and Split vs Floor-Standing form factors must never cross-match."""
    _assert_match(
        "Gree 1.5 Ton Inverter Split AC White",
        "Gree 1.5 Ton Inverter Floor Standing AC White",
        "Gree",
        False,
    )
    _assert_match(
        "Haier Inverter Split AC",
        "Haier HSU-18HFP Inverter Split AC",
        "Haier",
        False,
    )
    _assert_match(
        "Haier HSU-18HFCB Split AC 1.5 Ton",
        "Haier HPU-18HFCB Floor Standing Cabinet AC 1.5 Ton",
        "Haier",
        False,
    )


def _make_api_item(
    title: str,
    price: float,
    vendor_name: str,
    sold_out: bool = False,
) -> dict[str, Any]:
    """Construct a storefront API JSON item matching AlfaMall / DigiMall schema."""
    return {
        "title": title,
        "sold_out": sold_out,
        "variant_detail": {"discounted_price": price},
        "vendor": {"name": vendor_name, "handle": vendor_name.lower().replace(" ", "-")},
    }


def test_aud_crit_06_partial_waterfall_outage_no_false_monopoly() -> None:
    """AUD-CRIT-06: Own-store or sold-out-only Tier 1 results + Tier 2 HTTP 503 must flag SERVER_DOWN."""
    clear_scraper_cache()

    async def _run_outage_check() -> ScrapeResult:
        session = MagicMock()

        async def _fake_get(url: str, **kwargs: Any) -> _DummyResponse:
            params = kwargs.get("params") or {}
            if params.get("q") == "HSU-18HFPAA":
                # Tier 1 returns HTTP 200 with ONLY our own excluded store + a sold-out seller
                return _DummyResponse(
                    200,
                    {
                        "results": [
                            _make_api_item("Haier HSU-18HFPAA Split AC", 185000, "Kia Chahiye"),
                            _make_api_item(
                                "Haier HSU-18HFPAA Split AC",
                                179000,
                                "SoldOutVendor",
                                sold_out=True,
                            ),
                        ],
                        "total": 2,
                    },
                )
            # Tier 2 fails with HTTP 503 outage
            return _DummyResponse(503, {})

        session.get = AsyncMock(side_effect=_fake_get)
        with patch("pricing_engine.scraper.asyncio.sleep", new_callable=AsyncMock):
            return await scrape_alfamall(
                session,
                {"tier_1": ["HSU-18HFPAA"], "tier_2": ["18HFPAA"], "tier_3": []},
                target_model="HSU-18HFPAA",
                floor_price=160000.0,
                brand="Haier",
                target_desc="Haier Split AC HSU-18HFPAA",
            )

    scrape_res = asyncio.run(_run_outage_check())
    assert scrape_res.is_success is False
    assert len(scrape_res.own_offers) == 1

    floors = calculate_floors(cost_base=140000.0, min_profit=500.0)
    eval_res = evaluate_alfa_pricing(floors, scrape_res)
    assert eval_res.final_price is None
    assert "Server Down" in eval_res.status


# ==============================================================================
# HIGH-SEVERITY FINDINGS (AUD-HIGH-01 .. AUD-HIGH-10)
# ==============================================================================


def test_aud_high_01_modifier_source_end_to_end() -> None:
    """AUD-HIGH-01: Market-added modifiers must report source='market' instead of target omission."""
    clear_scraper_cache()

    async def _run() -> ScrapeResult:
        session = MagicMock()
        session.get = AsyncMock(
            return_value=_DummyResponse(
                200,
                {
                    "results": [
                        _make_api_item(
                            "Dawlance DWT 260 S Top Load Washing Machine",
                            78000,
                            "ExternalSeller",
                        )
                    ],
                    "total": 1,
                },
            )
        )
        return await scrape_alfamall(
            session,
            {"tier_1": ["DWT 260"], "tier_2": [], "tier_3": []},
            target_model="DWT 260",
            floor_price=65000.0,
            brand="Dawlance",
            target_desc="Dawlance DWT 260 Top Load Washing Machine",
        )

    res = asyncio.run(_run())
    assert "S" in (res.active_modifier_fallback or "")
    assert res.active_modifier_source == "market"
    assert res.offers[0].modifier_source == "market"

    floors = calculate_floors(cost_base=60000.0, min_profit=500.0)
    alfa_eval = evaluate_alfa_pricing(floors, res)
    assert "Marketplace listings specified" in alfa_eval.reason
    assert "(omitted in target)" in alfa_eval.reason
    assert "omitted in listings" not in alfa_eval.reason

    audit_txt, _, _ = format_fallback_descriptions(
        res.active_modifier_fallback or "S", source=res.active_modifier_source
    )
    assert "omitted in target" in audit_txt.lower() or "market" in audit_txt.lower()


def test_aud_high_02_winning_offer_modifier_attribution() -> None:
    """AUD-HIGH-02: Exact viable winners must not inherit Dual-Color or sub-floor modifiers."""
    floors = calculate_floors(cost_base=90000.0, min_profit=500.0)  # floor_alfa ~ 100,555.56

    # Case A: Exact single-color viable offer (108,000) beats dual-color offer (112,000)
    offers_a = ScrapeResult(
        offers=[
            ScrapedOffer(
                vendor_name="ExactSeller",
                price=108000.0,
                is_sold_out=False,
                raw_title="Gree GS-18PITH11W White",
                missing_modifier=None,
                modifier_source="target",
            ),
            ScrapedOffer(
                vendor_name="DualSeller",
                price=112000.0,
                is_sold_out=False,
                raw_title="Gree GS-18PITH11W White/Grey",
                missing_modifier="Dual-Color: White/Grey / Silver",
                modifier_source="dual_color",
            ),
        ],
        active_modifier_fallback=None,
    )
    res_a = evaluate_alfa_pricing(floors, offers_a)
    assert "CROSS-CHECK" not in res_a.reason
    assert "FALLBACK" not in res_a.reason

    # Case B: Sub-floor dumped offer (95,000) has Dual-Color, but actual undercut offer (109,000) is exact!
    offers_b = ScrapeResult(
        offers=[
            ScrapedOffer(
                vendor_name="SubFloorDual",
                price=95000.0,
                is_sold_out=False,
                raw_title="Gree GS-18PITH11W White/Grey",
                missing_modifier="Dual-Color: White/Grey / Silver",
                modifier_source="dual_color",
            ),
            ScrapedOffer(
                vendor_name="ViableExact",
                price=109000.0,
                is_sold_out=False,
                raw_title="Gree GS-18PITH11W White",
                missing_modifier=None,
                modifier_source="target",
            ),
        ],
        active_modifier_fallback=None,
    )
    res_b = evaluate_alfa_pricing(floors, offers_b)
    assert "Dual-Color" not in res_b.reason

    # Deduplication at identical (vendor, price) must preserve exact offer over dual-color duplicate
    deduped = deduplicate_offers(
        [
            ScrapedOffer(
                "SameVendor",
                108000.0,
                False,
                "Gree GS-18PITH11W White/Grey",
                missing_modifier="Dual-Color: White/Grey",
                modifier_source="dual_color",
            ),
            ScrapedOffer(
                "SameVendor",
                108000.0,
                False,
                "Gree GS-18PITH11W White",
                missing_modifier=None,
                modifier_source="target",
            ),
        ]
    )
    assert len(deduped) == 1
    assert deduped[0].missing_modifier is None


def test_aud_high_03_unit_preserving_capacity_norm() -> None:
    """AUD-HIGH-03: Distinct physical capacity units (8CFT vs 8KG, 1.5T vs 1.5KG) must collide."""
    _assert_match("Haier Hx-800 8 Cu Ft Fridge", "Haier Hx-800 8 KG Washer", "Haier", False)
    _assert_match("Pel P-15 1.5 Ton AC", "Pel P-15 1.5 KG Mini Appliance", "PEL", False)


def test_aud_high_04_pakistani_price_parsing_safety() -> None:
    """AUD-HIGH-04: parse_price_value must reject bool/NaN/Inf/<=0 and parse Pakistani currency formats."""
    assert parse_price_value(True) == 0.0
    assert parse_price_value(False) == 0.0
    assert parse_price_value(float("nan")) == 0.0
    assert parse_price_value(float("inf")) == 0.0
    assert parse_price_value(float("-inf")) == 0.0
    assert parse_price_value(-5000) == 0.0
    assert parse_price_value("NaN") == 0.0
    assert parse_price_value("Infinity") == 0.0

    assert parse_price_value("PKR: 75,000/-") == 75000.0
    assert parse_price_value("Rs: 82,500") == 82500.0
    assert parse_price_value("Rs. 124,999.00") == 124999.0
    assert parse_price_value("75,000 PKR") == 75000.0
    assert parse_price_value("75,000 Rs.") == 75000.0


def test_aud_high_05_washer_cross_facet_incompatibility() -> None:
    """AUD-HIGH-05: Front/Top Load automatic washers must never match Twin/Single Tub semi-auto washers."""
    _assert_match(
        "Haier HWM 85 Top Load Fully Automatic Washing Machine",
        "Haier HWM 85 Twin Tub Semi Automatic Washing Machine",
        "Haier",
        False,
    )
    _assert_match(
        "Dawlance DWF 8200 Front Load Washing Machine",
        "Dawlance DWF 8200 Single Tub Washer",
        "Dawlance",
        False,
    )
    # Regression guard: Dawlance 'DW-10600 W FL Twin Tub' must NOT self-collide on 'FL' lid code
    _assert_match(
        "Dawlance DW-10600 W FL Twin Tub Washing Machine",
        "Dawlance DW-10600 W FL Twin Tub Semi Automatic Washing Machine 10 KG",
        "Dawlance",
        True,
    )


def test_aud_high_06_generic_bundle_collision() -> None:
    """AUD-HIGH-06: Generic dual-SKU and Washer+Dryer/Spinner bundles must reject single-unit targets."""
    # Non-hardcoded brand prefix dual-SKU bundle (Super Asia SA-240 + SD-550)
    _assert_match(
        "Super Asia SA-240 Washing Machine",
        "Super Asia SA-240 Washer + SD-550 Spinner Combo Pack",
        "Super Asia",
        False,
    )
    # Washer & Dryer combo bundle vs single washer
    _assert_match(
        "Dawlance DW-7500 Washing Machine",
        "Dawlance DW-7500 Washer & Dryer Pair",
        "Dawlance",
        False,
    )
    # Single-unit Haier HWD Front Load Washer Dryer combo machine still matches itself
    _assert_match(
        "Haier HWD120-B14979U1 Front Load Washer Dryer",
        "Haier HWD120-B14979U1 12kg/8kg Washer & Dryer",
        "Haier",
        True,
    )


def test_aud_high_07_string_query_waterfall_parity() -> None:
    """AUD-HIGH-07: Passing a str query to scrape_alfamall/digimall must run waterfall promotion & telemetry."""
    clear_scraper_cache()

    async def _run() -> ScrapeResult:
        session = MagicMock()
        session.get = AsyncMock(
            return_value=_DummyResponse(
                200,
                {
                    "results": [
                        _make_api_item(
                            "Dawlance 9178LF Avante Refrigerator",
                            115000,
                            "MarketVendor",
                        )
                    ],
                    "total": 1,
                },
            )
        )
        return await scrape_digimall(
            session,
            "9178LF",
            target_model="9178LF GD",
            floor_price=95000.0,
            brand="Dawlance",
            target_desc="Dawlance 9178LF Glass Door Refrigerator",
        )

    res = asyncio.run(_run())
    assert res.executed_queries == ["9178LF"]
    assert len(res.offers) == 1
    assert res.active_modifier_fallback == "GD"


def test_aud_high_08_page2_transient_failure_no_cache_poison() -> None:
    """AUD-HIGH-08: A transient Page-2 HTTP 503 must not permanently mark p2_probed=True in cache."""
    clear_scraper_cache()
    p1_items = [
        _make_api_item("Gree GS-18PITH11W Split AC", 140000, "SubFloorSeller"),
    ] + [_make_api_item(f"Unrelated Model X-{i}", 50000 + i, f"V{i}") for i in range(17)]
    p2_attempts = 0

    async def _run_twice() -> tuple[ScrapeResult, ScrapeResult]:
        nonlocal p2_attempts
        session = MagicMock()

        async def _fake_get(url: str, **kwargs: Any) -> _DummyResponse:
            nonlocal p2_attempts
            params = kwargs.get("params") or {}
            if str(params.get("page", "1")) == "1":
                return _DummyResponse(200, {"results": p1_items, "total": 25})
            p2_attempts += 1
            if p2_attempts == 1:
                return _DummyResponse(503, {})
            return _DummyResponse(
                200,
                {
                    "results": [
                        _make_api_item(
                            "Gree GS-18PITH11W Split AC",
                            175000,
                            "Page2Winner",
                        )
                    ],
                    "total": 25,
                },
            )

        session.get = AsyncMock(side_effect=_fake_get)
        with patch("pricing_engine.scraper.asyncio.sleep", new_callable=AsyncMock):
            r1 = await scrape_alfamall(
                session,
                {"tier_1": ["GS-18PITH11W"], "tier_2": [], "tier_3": []},
                target_model="GS-18PITH11W",
                floor_price=150000.0,
                brand="Gree",
                target_desc="Gree GS-18PITH11W Split AC",
            )
            # Clear only high-level query cache (simulating second row sharing raw search query)
            _query_cache.clear()
            r2 = await scrape_alfamall(
                session,
                {"tier_1": ["GS-18PITH11W"], "tier_2": [], "tier_3": []},
                target_model="GS-18PITH11W",
                floor_price=150000.0,
                brand="Gree",
                target_desc="Gree GS-18PITH11W Split AC",
            )
        return r1, r2

    _, second_res = asyncio.run(_run_twice())
    assert any(o.vendor_name == "Page2Winner" for o in second_res.offers)


def test_aud_high_09_checkpoint_fingerprint_config_invalidation() -> None:
    """AUD-HIGH-09: Changing platform, personal_comm_pct, or Shipping Estimate must invalidate checkpoint rows."""
    row = pd.Series(
        {
            "Brand": "Gree",
            "Name": "Gree GS-18PITH11W",
            "Net Price": 150000,
            "Shipping Estimate": 3500,
        }
    )
    fp_alfa = compute_row_fingerprint(row, platform="alfa", personal_comm_pct=0.0)
    fp_all = compute_row_fingerprint(row, platform="all", personal_comm_pct=0.0)
    fp_comm = compute_row_fingerprint(row, platform="all", personal_comm_pct=5.0)
    row_ship = row.copy()
    row_ship["Shipping Estimate"] = 4500
    fp_ship = compute_row_fingerprint(row_ship, platform="all", personal_comm_pct=0.0)

    assert len({fp_alfa, fp_all, fp_comm, fp_ship}) == 4

    df = pd.DataFrame([row])
    state = CheckpointState(
        catalog_name="c.csv",
        catalog_path="c.csv",
        catalog_mtime=1.0,
        total_items=1,
        completed_records={0: {"S.NO": 1, "Final Price Alfa": 170000}},
        row_fingerprints={0: fp_alfa},
    )
    rows_to_run, cached = detect_dirty_or_pending_rows(
        df, state, platform="all", personal_comm_pct=0.0
    )
    assert rows_to_run == [0]
    assert cached == {}


def test_aud_high_10_ci_workflow_gates() -> None:
    """AUD-HIGH-10: CI workflow must enforce ruff check, ruff format --check, mypy, and pytest."""
    repo_root = Path(__file__).resolve().parent.parent
    ci_text = (repo_root / ".github" / "workflows" / "ci.yml").read_text(encoding="utf-8")
    assert "ruff check ." in ci_text
    assert "ruff format --check ." in ci_text
    assert "mypy ." in ci_text
    assert "pytest" in ci_text

    pyproject_text = (repo_root / "pyproject.toml").read_text(encoding="utf-8")
    assert '".agents"' in pyproject_text


# ==============================================================================
# MEDIUM-SEVERITY FINDINGS (AUD-MED-01 .. AUD-MED-07)
# ==============================================================================


def test_aud_med_01_csv_xlsx_parity_and_dynamic_comm_labels(tmp_path: Path) -> None:
    """AUD-MED-01: CSV & Excel math must agree to the rupee, and -m <pct> labels must be dynamic."""
    # 1. Volumetric shipping 2-decimal ceiling parity: (33*33*33)/5000 * 70 = 503.118 -> 503.12
    vw = calculate_volumetric_weight(33, 33, 33)
    ship = calculate_shipping_estimate(vw)
    assert ship == 503.12

    # 2. Dynamic 8% personal commission run (-m 8)
    product = parse_product_input(
        0,
        pd.Series(
            {
                "S.NO": 1,
                "Brand": "Gree",
                "Name": "Gree GS-18PITH11W",
                "Length": 100,
                "Width": 40,
                "Height": 90,
                "Net Price": 130000,
            }
        ),
    )
    offers = [ScrapedOffer("CompA", 165000.0, False, "Gree GS-18PITH11W")]
    rec = process_product_with_offers(product, offers, offers, personal_comm_pct=8.0)

    assert "My 8% Commission" in rec
    assert "Min Profit Target (8% + 500)" in rec
    assert "Alfa 8% Goal" in rec
    assert "Digi 8% Goal" in rec
    assert rec["Alfa 8% Goal"] == "WIN + 8% ACHIEVED"
    assert rec["Min Profit Target (8% + 500)"] == calculate_effective_target_profit(
        personal_comm_pct=8.0
    )

    df = pd.DataFrame([rec])
    manifest = export_results(
        df=df,
        source_filename="comm8_test.csv",
        output_dir=tmp_path / "full",
        render_panel=False,
        personal_comm_pct=8.0,
    )
    assert manifest.latest_xlsx_path is not None
    wb = openpyxl.load_workbook(manifest.latest_xlsx_path)
    ws = wb.active
    assert ws is not None
    headers = [cell.value for cell in ws[1]]
    col_map = {h: openpyxl.utils.get_column_letter(i + 1) for i, h in enumerate(headers)}

    ship_formula = str(ws[f"{col_map['Shipping Estimate']}2"].value)
    assert ship_formula.startswith("=ROUNDUP(") and ship_formula.endswith(", 2)")

    min_target_formula = str(ws[f"{col_map['Min Profit Target (8% + 500)']}2"].value)
    assert min_target_formula == "=ROUNDUP(500/(1-0.08), 0)"

    my_comm_formula = str(ws[f"{col_map['My 8% Commission']}2"].value)
    assert "*0.08" in my_comm_formula

    w_csv, w_xlsx, w_count = export_winning_5pct_results(
        df=df,
        source_filename="comm8_test.csv",
        output_dir=tmp_path / "winners",
        personal_comm_pct=8.0,
    )
    assert w_count == 1
    assert w_csv is not None and "winning_8pct" in w_csv.name
    assert w_xlsx is not None and "winning_8pct" in w_xlsx.name


def test_aud_med_02_multicolor_slash_and_crosscheck_badge() -> None:
    """AUD-MED-02: 3-way color slashes must expand cleanly and trigger Dual-Color CROSS-CHECK badges."""
    branches = expand_cartesian_branches("Haier HSU-18HFP White / Silver / Black Split AC")
    assert len(branches) == 3
    for color_word in ("White", "Silver", "Black"):
        matching = [b for b in branches if color_word in b]
        assert len(matching) == 1
        assert "/" not in matching[0]

    # Target without explicit color vs candidate with slash color chain triggers Dual-Color cross-check
    t_slots = extract_product_slots("Dawlance 9178LF Refrigerator", brand="Dawlance")
    dec = referee_match(
        t_slots, "Dawlance 9178LF Chrome / Silver / Black Refrigerator", brand="Dawlance"
    )
    assert dec.is_match is True
    assert dec.modifier_source == "market"
    assert dec.missing_modifier is not None and dec.missing_modifier.startswith("Dual-Color")

    # Executive summary price formatter renders CROSS-CHECK on '[CROSS-CHECK NOTICE: ...]' reasons
    summary_cell = _format_summary_price(
        114899,
        min_comp_price=115000,
        reason="Competitive Win [CROSS-CHECK NOTICE: Dual-Color listing matched]",
        alert="NONE",
    )
    assert "CROSS-CHECK" in summary_cell


def test_aud_med_03_iot_crosscheck_and_multimodifier_join() -> None:
    """AUD-MED-03: Extra candidate IOT/WIFI on numeric stems requires cross-check; multi-modifiers join."""
    # Candidate adds IOT on a pure-numeric refrigerator stem -> base match with market cross-check
    t_iot = extract_product_slots("Haier HRF-458 Refrigerator", brand="Haier")
    dec_iot = referee_match(t_iot, "Haier HRF-458 IOT Smart Refrigerator", brand="Haier")
    assert dec_iot.is_match is False
    assert dec_iot.is_base_match is True
    assert dec_iot.modifier_source == "market"
    assert dec_iot.missing_modifier is not None and "IOT" in dec_iot.missing_modifier

    # Target has both Glass Door (GD) and Inverter, candidate omits both -> both reported in base match
    t_multi = extract_product_slots(
        "Dawlance 9178LF Glass Door Inverter Refrigerator", brand="Dawlance"
    )
    dec_multi = referee_match(t_multi, "Dawlance 9178LF Refrigerator", brand="Dawlance")
    assert dec_multi.is_match is False
    assert dec_multi.is_base_match is True
    assert dec_multi.missing_modifier is not None
    assert "GD" in dec_multi.missing_modifier
    assert "INVERTER" in dec_multi.missing_modifier


def test_aud_med_04_supabase_param_encoding_and_session_close() -> None:
    """AUD-MED-04: SupabaseClient must URL-encode PostgREST filter values and close its AsyncSession."""
    client = SupabaseClient(url="https://example.supabase.co", key="test-key")
    mock_sess = MagicMock()
    mock_sess.get = AsyncMock(return_value=_DummyResponse(200, []))
    mock_sess.close = AsyncMock()
    client._session = mock_sess

    async def _exercise() -> None:
        await client.fetch_approved_capacities("Black & Decker")
        await client.fetch_approved_synonyms("Super Asia", "Washer & Dryer")
        await client.close()

    asyncio.run(_exercise())

    called_urls = [c.args[0] for c in mock_sess.get.call_args_list]
    assert any("BLACK%20%26%20DECKER" in u for u in called_urls)
    assert any("SUPER%20ASIA" in u and "WASHER%20%26%20DRYER" in u for u in called_urls)
    mock_sess.close.assert_awaited_once()
    assert client._session is None


def test_aud_med_05_pipe_noise_and_capacity_registry_ssot() -> None:
    """AUD-MED-05: extract_search_query_and_model strips pipe seller noise & capacity uses InMemoryRegistry."""
    query, model = extract_search_query_and_model(
        "Gree GS-18PITH11W Split AC | Official Warranty | Best Price in Pakistan",
        brand="Gree",
    )
    assert "PAKISTAN" not in model.upper()
    assert "WARRANTY" not in model.upper()
    assert "18PITH11W" in model.upper()
    assert "18PITH11W" in query.upper()

    # Single source of truth: dynamic InMemoryRegistry injection reflects immediately in get_capacity_tonnage
    InMemoryRegistry._ensure_initialized()
    orig_map = dict(InMemoryRegistry._capacity_cache)
    try:
        InMemoryRegistry._capacity_cache[("TESTBRAND", "19")] = "1.6"
        assert get_capacity_map("TESTBRAND")["19"] == "1.6"
        assert get_capacity_tonnage("19", brand="TESTBRAND") == "1.6"
    finally:
        InMemoryRegistry._capacity_cache = orig_map


def test_aud_med_06_run_scoped_cache_and_learner_cleanup(tmp_path: Path) -> None:
    """AUD-MED-06: Pipeline runs must clear stale scraper caches on entry and learner buffers on exit."""
    _query_cache[("alfa", "STALE", "STALE")] = ScrapeResult(
        offers=[ScrapedOffer("StaleV", 99999.0, False, "Stale")]
    )
    _raw_items_cache[("alfa", "STALE")] = [{"title": "Stale"}]
    _raw_items_meta_cache[("alfa", "STALE")] = (1, True)
    default_learner.buffer_observation(
        brand="Haier",
        category="AC",
        anchor_model="18HFP",
        target_residual="INV",
        candidate_residual="DC",
        vendor_name="V1",
        raw_title="Haier 18HFP DC",
    )
    assert default_learner.buffer_size == 1

    cat = tmp_path / "mini.csv"
    cat.write_text("Brand,Name,Net Price\nGree,GS-18PITH11W,140000\n", encoding="utf-8")

    with (
        patch(
            "pricing_engine.pipeline.scrape_alfamall",
            new=AsyncMock(return_value=ScrapeResult(offers=[])),
        ),
        patch(
            "pricing_engine.pipeline.scrape_digimall",
            new=AsyncMock(return_value=ScrapeResult(offers=[])),
        ),
    ):
        asyncio.run(
            run_pipeline(
                input_path=cat,
                export_dir=tmp_path / "out",
                verbose=False,
                resume=False,
            )
        )

    assert ("alfa", "STALE") not in _raw_items_cache
    assert ("alfa", "STALE") not in _raw_items_meta_cache
    assert default_learner.buffer_size == 0


def test_aud_med_07_range_pipeline_brand_validation(tmp_path: Path) -> None:
    """AUD-MED-07: Range mode must preserve multi-word brands in Name-only CSVs and reject invalid brands."""
    # Multi-word & connector brands in Name-only CSV
    name_only = tmp_path / "names.csv"
    name_only.write_text(
        "Name\n"
        "Super Asia SA-240 Washing Machine\n"
        "Super Asia ECM-4500 Plus Air Cooler\n"
        "Black & Decker BX-200 Blender\n"
        "Haier PuriCool HSU-18HFP Split AC\n"
        "Dawlance Avante 9178LF Refrigerator\n",
        encoding="utf-8",
    )
    products = load_range_products(name_only)
    assert [p.brand for p in products] == [
        "Super Asia",
        "Super Asia",
        "Black & Decker",
        "Haier",
        "Dawlance",
    ]

    # Placeholder or numeric-only explicit Brand must raise ValueError
    for bad_brand in ("N/A", "Unknown", "12345", "---"):
        bad_csv = tmp_path / f"bad_{abs(hash(bad_brand))}.csv"
        bad_csv.write_text(f"Brand,Name\n{bad_brand},GS-18PITH11W\n", encoding="utf-8")
        with pytest.raises(ValueError, match="Brand"):
            load_range_products(bad_csv)

    # Category-only Name without brand in Name-only CSV must also raise ValueError
    nobrand_csv = tmp_path / "nobrand.csv"
    nobrand_csv.write_text("Name\nSplit AC 1.5 Ton Inverter\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Brand"):
        load_range_products(nobrand_csv)


# ==============================================================================
# LOW-SEVERITY & NEWLY DISCOVERED FINDINGS (AUD-LOW-01..04, AUD-NEW-01..03)
# ==============================================================================


def test_aud_low_01_learner_metric_unit_guards() -> None:
    """AUD-LOW-01: AutonomousLearner must block metric units (CFT, TAP, BTU, WATT, VOLT, AMP) & Greek TON."""
    learner = AutonomousLearner()
    for bad_unit in ("CFT", "TAP", "BTU", "WATT", "VOLT", "AMP", "\u03a4\u039fN"):
        learner.buffer_observation(
            brand="Dawlance",
            category="REFRIGERATOR",
            anchor_model="9178",
            target_residual=bad_unit,
            candidate_residual="LVS",
            vendor_name="Seller1",
            raw_title="Dawlance 9178",
        )
    assert learner.buffer_size == 0


def test_aud_low_02_gitignore_repomix_glob() -> None:
    """AUD-LOW-02: .gitignore must ignore repomix-output.* dumps."""
    repo_root = Path(__file__).resolve().parent.parent
    gitignore_text = (repo_root / ".gitignore").read_text(encoding="utf-8")
    assert "repomix-output.*" in gitignore_text


def test_aud_low_04_range_stagger_consistency(tmp_path: Path) -> None:
    """AUD-LOW-04: run_range_pipeline must stagger all N-1 item launches, not only concurrency-1."""
    cat = tmp_path / "five.csv"
    cat.write_text(
        "Brand,Name\nGree,GS-12A\nGree,GS-12B\nGree,GS-12C\nGree,GS-12D\nGree,GS-12E\n",
        encoding="utf-8",
    )
    sleep_mock = AsyncMock()
    with (
        patch(
            "pricing_engine.range_pipeline.scrape_alfamall",
            new=AsyncMock(return_value=ScrapeResult(offers=[])),
        ),
        patch(
            "pricing_engine.range_pipeline.scrape_digimall",
            new=AsyncMock(return_value=ScrapeResult(offers=[])),
        ),
        patch("pricing_engine.range_pipeline.RANGE_STAGGER_SECONDS", 0.01),
        patch("pricing_engine.range_pipeline.asyncio.sleep", sleep_mock),
    ):
        asyncio.run(
            run_range_pipeline(cat, concurrency=2, verbose=False, export_dir=tmp_path / "out")
        )
    assert sleep_mock.await_count == 4  # N - 1 == 5 - 1 == 4


def test_aud_new_01_platform_personal_comm_and_config_ssot() -> None:
    """AUD-NEW-01 & AUD-LOW-03: Platform evaluators must honor distinct Alfa vs Digi personal_comm & config."""
    cost_base = calculate_cost_base(100000.0, 2000.0)
    with (
        patch.object(config, "ALFA_PERSONAL_COMM_PCT", 4.0),
        patch.object(config, "DIGI_PERSONAL_COMM_PCT", 8.0),
    ):
        floors = calculate_floors(
            cost_base=cost_base,
            min_profit=500.0,
            personal_comm_pct=config.DEFAULT_PERSONAL_COMM_PCT,
        )
    assert floors.alfa_personal_comm != floors.digi_personal_comm

    # All competitors dump below floor -> Safe Fallback uses platform-specific personal_comm
    dumped = [ScrapedOffer("LossLeader", 90000.0, False, "Model X")]
    alfa_res = evaluate_alfa_pricing(floors, dumped)
    digi_res = evaluate_digi_pricing(floors, dumped)
    assert alfa_res.final_price is not None and alfa_res.real_profit is not None
    assert digi_res.final_price is not None and digi_res.real_profit is not None
    assert alfa_res.real_profit >= 1500 + floors.alfa_personal_comm - 1
    assert digi_res.real_profit >= 1500 + floors.digi_personal_comm - 1


def test_aud_new_02_digi_fractional_floor_band_handling() -> None:
    """AUD-NEW-02: Competitor inside (floor_digi_8, ceil(floor_digi_8)] must drop to upper-band undercut."""
    # cost_base = 53500 -> floor_digi_8 = (53500 + 500) / 0.92 = 58695.65217... (ceil = 58696)
    floors = calculate_floors(cost_base=53500.0, min_profit=500.0)
    assert 58695.65 < floors.floor_digi_8 < 58696.0

    # Offer 1 at 58696.0 cannot be strictly undercut by any integer >= 58696; Offer 2 at 60500.0 is viable!
    offers = [
        ScrapedOffer("AtCeilBoundary", 58696.0, False, "Model Z"),
        ScrapedOffer("ViableUpper", 60500.0, False, "Model Z"),
    ]
    res = evaluate_digi_pricing(floors, offers)
    assert res.final_price is not None
    assert math.ceil(floors.floor_digi_8) <= res.final_price < 60500
    assert len(res.dropped) == 1
    assert "AtCeilBoundary" in res.dropped[0]


def test_aud_new_03_ambiguity_resolution_preserves_scrape_metadata() -> None:
    """AUD-NEW-03: Selecting a variant during interactive ambiguity conflict must preserve ScrapeResult metadata."""
    product = ProductInput(
        sno=1,
        brand="Dawlance",
        name="Dawlance DW-10600",
        length=50,
        width=50,
        height=90,
        net_price=40000.0,
        profit=500.0,
        raw_row={"S.NO": 1, "Brand": "Dawlance", "Name": "Dawlance DW-10600", "Net Price": 40000},
    )
    off_c = ScrapedOffer("Comp1", 55000.0, False, "Dawlance DW-10600 C", detected_variant="C")
    off_g = ScrapedOffer("Comp2", 56000.0, False, "Dawlance DW-10600 G", detected_variant="G")
    own_off = ScrapedOffer("Kia Chahiye", 54500.0, False, "Dawlance DW-10600 C")
    conflict = AmbiguousVariantConflict(
        stem_model="DW-10600",
        detected_variants=["C", "G"],
        variant_offers={"C": [off_c], "G": [off_g]},
    )
    scraped = ScrapeResult(
        offers=[off_c, off_g],
        own_offers=[own_off],
        executed_queries=["DW-10600"],
        matched_tier="tier_1",
        tier_summary={"tier_1": 2},
        ambiguity_conflict=conflict,
    )

    captured_alfa: list[Any] = []
    real_eval = evaluate_alfa_pricing

    def _spy_eval(floors: Any, offers: Any) -> Any:
        captured_alfa.append(offers)
        return real_eval(floors, offers)

    with (
        patch("pricing_engine.pipeline.scrape_alfamall", new=AsyncMock(return_value=scraped)),
        patch(
            "pricing_engine.pipeline.scrape_digimall",
            new=AsyncMock(return_value=ScrapeResult(offers=[])),
        ),
        patch("pricing_engine.pipeline.evaluate_alfa_pricing", side_effect=_spy_eval),
        patch("pricing_engine.pipeline.sys.stdin.isatty", return_value=True),
        patch("builtins.input", return_value="1"),
    ):
        rec = asyncio.run(
            process_item_async(
                product=product,
                session=MagicMock(),
                semaphore=asyncio.Semaphore(1),
                verbose=False,
            )
        )

    assert "DW-10600 C" in rec["Name"]
    assert len(captured_alfa) == 1
    preserved = captured_alfa[0]
    assert isinstance(preserved, ScrapeResult)
    assert preserved.own_offers == [own_off]
    assert preserved.executed_queries == ["DW-10600"]
    assert preserved.tier_summary == {"tier_1": 2}
    assert preserved.ambiguity_conflict is None


def test_all_13_pakistani_brands_isolation_and_own_store_exclusions() -> None:
    """Verify exact match & cross-variant rejection across all 13 Pakistani brands + own-store exclusions."""
    brand_cases = [
        (
            "Haier",
            "Haier HSU-18HFP Split AC",
            "Haier HSU-18HFP Inverter Split AC",
            "Haier HSU-24HFP Split AC",
        ),
        (
            "Dawlance",
            "Dawlance 9178LF WB Fridge",
            "Dawlance 9178LF WB Refrigerator",
            "Dawlance 9178LF GB Refrigerator",
        ),
        (
            "PEL",
            "PEL PRGD-2350 Glass Door Fridge",
            "PEL PRGD-2350 Glass Door Refrigerator",
            "PEL PRGD-2550 Glass Door Refrigerator",
        ),
        (
            "Orient",
            "Orient Ultron 18G Split AC",
            "Orient Ultron 18G DC Inverter AC",
            "Orient Ultron 24G DC Inverter AC",
        ),
        (
            "Kenwood",
            "Kenwood KEE-1844S Split AC",
            "Kenwood KEE-1844S eSmart Inverter AC",
            "Kenwood KEE-1244S eSmart AC",
        ),
        (
            "Gree",
            "Gree GS-18PITH11W Split AC",
            "Gree GS-18PITH11W 1.5 Ton Split AC",
            "Gree GS-18PITH11S 1.5 Ton Split AC",
        ),
        (
            "Midea",
            "MideaMSAG-18CRDN8 Split AC".replace("MideaM", "Midea M"),
            "Midea MSAG-18CRDN8 1.5 Ton Inverter AC",
            "Midea MSAG-24CRDN8 2 Ton AC",
        ),
        (
            "TCL",
            "TCL TAC-18T3-PRO Split AC",
            "TCL TAC-18T3-PRO 1.5 Ton T3 Inverter AC",
            "TCL TAC-24T3-PRO 2 Ton AC",
        ),
        (
            "Hisense",
            "Hisense AS-18TR4SVETG01 Split AC",
            "Hisense AS-18TR4SVETG01 1.5 Ton AC",
            "Hisense AS-24TR4SVETG01 2 Ton AC",
        ),
        (
            "EcoStar",
            "EcoStar CX-43UD962 43 Inch 4K LED TV",
            "EcoStar CX-43UD962 UHD Smart LED TV",
            "EcoStar CX-50UD962 50 Inch LED TV",
        ),
        (
            "Homage",
            "Homage HWD-49332P Water Dispenser",
            "Homage HWD-49332P 3 Taps Water Dispenser",
            "Homage HWD-49331P Water Dispenser",
        ),
        (
            "Super Asia",
            "Super Asia SA-240 Washing Machine",
            "Super Asia SA-240 Single Tub Washer",
            "Super Asia SD-240 Spinner Dryer",
        ),
        (
            "Westpoint",
            "Westpoint WF-5256 Air Fryer",
            "Westpoint WF-5256 Digital Air Fryer",
            "Westpoint WF-5258 Digital Air Fryer",
        ),
    ]
    for brand, target, good_cand, bad_cand in brand_cases:
        _assert_match(target, good_cand, brand, True)
        _assert_match(target, bad_cand, brand, False)

    # Own-store & official flagship exclusions
    for excluded_vendor in (
        "Kia Chahiye",
        "kiachahye.com",
        "Kia Chahiye Pvt Ltd",
        "Ajmer Electronics",
        "ajmer-electronics",
        "Haier Official Store",
        "haier-official-store",
    ):
        assert is_non_competitor_vendor(excluded_vendor) is True
    assert is_non_competitor_vendor("Electro Gas") is False
