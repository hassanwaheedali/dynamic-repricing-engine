"""Unit and integration tests for default category shipping fallback logic."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from pricing_engine.calculator import resolve_fallback_shipping
from pricing_engine.exporter import export_results
from pricing_engine.ingestion import ingest_catalog
from pricing_engine.models import CostFloors, PlatformPricingResult, ProductInput
from pricing_engine.pipeline import (
    console,
    render_audit_card,
    resolve_shipping_estimate,
)


class TestResolveFallbackShipping:
    """Test category, sub-type, and capacity detection in resolve_fallback_shipping."""

    # 1. Single Tub Washers
    @pytest.mark.parametrize(
        ("name", "expected_rate"),
        [
            ("Haier Single Tub Washing Machine 8 KG HWM-80", 2500.0),
            ("Dawlance Single Tub Washer 7.5kg", 2500.0),
            ("Super Asia Single Tub 10 kg Washer Only", 3000.0),
            ("Kenwood Single Tub 12 kg Washing Machine", 3500.0),
            ("Haier Single Tub Washer 15 KG", 3500.0),
        ],
    )
    def test_single_tub_brackets(self, name: str, expected_rate: float) -> None:
        rate, label = resolve_fallback_shipping(name, brand="Haier")
        assert rate == expected_rate
        assert "Single Tub" in label

    # 2. Twin Tub Washers
    @pytest.mark.parametrize(
        ("name", "expected_rate"),
        [
            ("Haier Twin Tub Washing Machine 8 KG", 3000.0),
            ("Dawlance Semi-Automatic Twin Tub 8.5 kg", 3000.0),
            ("Super Asia Twin Tub 10 kg Washer & Spinner", 3500.0),
            ("Haier Twin Tub 12 KG Machine", 3800.0),
            ("Dawlance Twin Tub 15 KG Jumbo Washer", 4500.0),
        ],
    )
    def test_twin_tub_brackets(self, name: str, expected_rate: float) -> None:
        rate, label = resolve_fallback_shipping(name, brand="Haier")
        assert rate == expected_rate
        assert "Twin Tub" in label

    # 3. Automatic Washers
    @pytest.mark.parametrize(
        ("name", "expected_rate"),
        [
            ("Haier Fully Automatic Top Load Washing Machine 8.5 KG", 3500.0),
            ("Dawlance Automatic Front Load Washer 9.5 kg", 3500.0),
            ("Haier Automatic Washing Machine 10 kg HWM 100", 3500.0),
            ("Super Asia Fully Automatic 12 kg Washer", 4500.0),
            ("Dawlance Automatic Inverter Washer 15 KG", 4500.0),
        ],
    )
    def test_automatic_brackets(self, name: str, expected_rate: float) -> None:
        rate, label = resolve_fallback_shipping(name, brand="Haier")
        assert rate == expected_rate
        assert "Automatic" in label

    # 4. Model number capacity extraction (HWM-120 -> 12kg, HW 105 -> 10.5kg, etc.)
    def test_washer_model_number_capacity_extraction(self) -> None:
        rate, label = resolve_fallback_shipping("Haier Fully Automatic HWM-120-1678", brand="Haier")
        assert rate == 4500.0
        assert "12kg" in label

        rate, label = resolve_fallback_shipping("Haier Twin Tub HWM 80-826", brand="Haier")
        assert rate == 3000.0
        assert "8kg" in label

    # 5. Air Conditioners
    @pytest.mark.parametrize(
        ("name", "expected_rate", "expected_ton"),
        [
            ("Gree Inverter AC 1.0 Ton 12AITH21W", 3000.0, "1.0 Ton"),
            ("Haier Split AC 12000 BTU 1 Ton", 3000.0, "1.0 Ton"),
            ("Gree Inverter AC 1.5 Ton 18AITH21W", 3500.0, "1.5 Ton"),
            ("Kenwood eInverter AC 18000 BTU", 3500.0, "1.5 Ton"),
            ("Gree Inverter AC 2.0 Ton 24AITH21W", 4500.0, "2.0 Ton"),
            ("Haier Inverter AC 24000 BTU 2 Ton", 4500.0, "2.0 Ton"),
            ("Dawlance Split AC Inverter", 3500.0, "Air Conditioner"),
        ],
    )
    def test_air_conditioner_fallback(self, name: str, expected_rate: float, expected_ton: str) -> None:
        rate, label = resolve_fallback_shipping(name, brand="Gree")
        assert rate == expected_rate
        assert expected_ton in label

    # 6. Refrigerators
    @pytest.mark.parametrize(
        ("name", "expected_rate", "expected_sub"),
        [
            ("Dawlance 4467 SD Mini Bar Refrigerator", 4000.0, "Small / Mini Bar"),
            ("Haier Deep Freezer 500L Side by Side", 7500.0, "Large / Side-by-Side"),
            ("Dawlance 9170 Inverter Refrigerator Glass Door", 5500.0, "Refrigerator"),
        ],
    )
    def test_refrigerator_fallback(self, name: str, expected_rate: float, expected_sub: str) -> None:
        rate, label = resolve_fallback_shipping(name, brand="Dawlance")
        assert rate == expected_rate
        assert expected_sub in label

    # 7. General Appliance Default
    def test_general_default(self) -> None:
        rate, label = resolve_fallback_shipping("Philips Electric Kettle", brand="Philips")
        assert rate == 3000.0
        assert label == "General Appliance"

    # 8. Dynamic Custom Category Extensibility (Zero Code Changes in calculator.py)
    def test_user_defined_custom_category_without_code_changes(self) -> None:
        custom_config = {
            "categories": [
                {
                    "id": "microwave",
                    "name": "Microwave Oven",
                    "keywords": ["microwave", "oven", "mwo"],
                    "default": 1500.0,
                },
                {
                    "id": "led_tv",
                    "name": "LED TV",
                    "keywords": ["led", "tv", "smart tv"],
                    "sub_types": [
                        {"name": "LED TV (32-43 Inch)", "pattern": r"\b(?:32|40|42|43)\b", "rate": 2000.0},
                        {"name": "LED TV (50-65 Inch)", "pattern": r"\b(?:50|55|65)\b", "rate": 3500.0},
                    ],
                    "default": 2500.0,
                },
            ],
            "default": 3000.0,
        }

        # Flat category resolution
        rate, label = resolve_fallback_shipping("Dawlance Microwave Oven 20L", brand="Dawlance", config_table=custom_config)
        assert rate == 1500.0
        assert label == "Microwave Oven"

        # Bracket category resolution
        rate, label = resolve_fallback_shipping("TCL 43 Inch 4K Smart Google TV LED", brand="TCL", config_table=custom_config)
        assert rate == 2000.0
        assert label == "LED TV (32-43 Inch)"

        rate, label = resolve_fallback_shipping("Samsung 55 Inch QLED 4K TV", brand="Samsung", config_table=custom_config)
        assert rate == 3500.0
        assert label == "LED TV (50-65 Inch)"


class TestResolveShippingEstimate:
    """Test resolution priority: CSV override > LxWxH volumetric > category fallback."""

    def test_csv_shipping_override_takes_precedence(self) -> None:
        prod = ProductInput(
            sno=1,
            brand="Haier",
            name="Haier Automatic 10kg",
            length=60,
            width=60,
            height=85,
            net_price=70000,
            profit=500,
            raw_row={"Shipping Estimate": "4200"},
        )
        shipping, is_fallback = resolve_shipping_estimate(prod)
        assert shipping == 4200.0
        assert is_fallback is False

    def test_volumetric_dimensions_used_when_positive(self) -> None:
        prod = ProductInput(
            sno=1,
            brand="Haier",
            name="Haier Automatic 10kg",
            length=60,
            width=60,
            height=85,
            net_price=70000,
            profit=500,
            raw_row={},
        )
        # Volumetric: (60 * 60 * 85) / 5000 = 61.2 kg * 70 = 4284.0
        shipping, is_fallback = resolve_shipping_estimate(prod)
        assert shipping == 4284.0
        assert is_fallback is False

    def test_missing_dimensions_triggers_fallback(self) -> None:
        prod = ProductInput(
            sno=1,
            brand="Haier",
            name="Haier Fully Automatic Top Load 10 KG",
            length=0,
            width=0,
            height=0,
            net_price=70000,
            profit=500,
            raw_row={},
        )
        shipping, is_fallback = resolve_shipping_estimate(prod)
        assert shipping == 3500.0
        assert is_fallback is True


class TestCatalogIngestionWithoutDimensions:
    """Test that CSVs without Length, Width, Height ingest smoothly without errors."""

    def test_ingest_catalog_3_columns_only(self, tmp_path: Path) -> None:
        csv_file = tmp_path / "minimal_catalog.csv"
        csv_file.write_text(
            "Brand,Name,Net Price\n"
            "Haier,Haier Automatic Washing Machine 12 KG,75000\n"
            "Gree,GREE SPLIT AC 1.5 TON 18AITH21W,185000\n",
            encoding="utf-8",
        )
        products, report = ingest_catalog(csv_file, render_banner=False)
        assert len(products) == 2
        assert report.valid_rows == 2

        # Check product 1 (Washer 12kg automatic)
        p1 = products[0]
        assert p1.length == 0.0
        assert p1.width == 0.0
        assert p1.height == 0.0
        ship1, is_fb1 = resolve_shipping_estimate(p1)
        assert ship1 == 4500.0
        assert is_fb1 is True

        # Check product 2 (AC 1.5 Ton)
        p2 = products[1]
        ship2, is_fb2 = resolve_shipping_estimate(p2)
        assert ship2 == 3500.0
        assert is_fb2 is True


class TestTerminalAuditCardFormatting:
    """Test terminal audit card formatting when shipping fallback is active."""

    def test_render_audit_card_shows_fallback_badge(self) -> None:
        product = ProductInput(
            sno=1,
            brand="Haier",
            name="Haier Washing Machine Twin Tub 8 KG",
            length=0.0,
            width=0.0,
            height=0.0,
            net_price=45000.0,
            profit=500.0,
            raw_row={},
        )
        floors = CostFloors(
            cost_base=48000.0,
            zero_margin_10=53333.33,
            zero_margin_8=52173.91,
            floor_alfa=53888.89,
            floor_digi_10=53888.89,
            floor_digi_8=52717.39,
            personal_comm=0,
        )

        mock_alfa_res = PlatformPricingResult(
            min_price=55000.0,
            max_price=60000.0,
            final_price=54999,
            real_profit=1499,
            status="SUCCESS",
            reason="COMPETITIVE_WIN",
            effective_comm="10.00%",
        )

        mock_digi_res = PlatformPricingResult(
            min_price=55000.0,
            max_price=60000.0,
            final_price=54999,
            real_profit=2599,
            status="SUCCESS",
            reason="COMPETITIVE_WIN",
            effective_comm="8.00%",
        )

        # Render with is_shipping_fallback = True with console capture
        with console.capture() as capture:
            render_audit_card(
                product=product,
                shipping=3000.0,
                floors=floors,
                baseline_calc=55000.0,
                target_model="HWM80",
                alfa_offers=[],
                digi_offers=[],
                alfa_res=mock_alfa_res,
                digi_res=mock_digi_res,
                is_shipping_fallback=True,
            )

        output = capture.get()
        assert "Default Config Shipping" in output
        assert "Not Provided" in output


class TestExcelExportFallbackFormula:
    """Test that Excel export retains fallback shipping in formula when dimensions are missing."""

    def test_excel_formula_contains_fallback_guard(self, tmp_path: Path) -> None:
        df = pd.DataFrame(
            [
                {
                    "S.NO": 1,
                    "Brand": "Haier",
                    "Name": "Haier Washer with Dimensions",
                    "Net Price": 70000,
                    "Length": 60,
                    "Width": 60,
                    "Height": 85,
                    "Profit": 500,
                    "Shipping Estimate": 4284,
                    "Commision": "10.00%",
                    "Digi Effective Commision": "10.00%",
                    "Total Without Profit": 82538,
                    "Calculation": 83093,
                    "Minimum Price Alfa": 83093,
                    "Maximum Price Alfa": 83093,
                    "Minimum Price Digi": 83093,
                    "Maximum Price Digi": 83093,
                    "Final Price Alfa": 83999,
                    "Final Price Digi": 83999,
                    "Alfa Real Profit": 1315,
                    "Digi Real Profit": 1315,
                    "Alfa Reason": "MONOPOLY_PRICING",
                    "Digi Reason": "MONOPOLY_PRICING",
                    "Alfa Risk Alert": "NONE",
                    "Digi Risk Alert": "NONE",
                },
                {
                    "S.NO": 2,
                    "Brand": "Haier",
                    "Name": "Haier Automatic 10kg Missing Dimensions",
                    "Net Price": 70000,
                    "Length": 0,
                    "Width": 0,
                    "Height": 0,
                    "Profit": 500,
                    "Shipping Estimate": 3500,
                    "Commision": "10.00%",
                    "Digi Effective Commision": "10.00%",
                    "Total Without Profit": 81667,
                    "Calculation": 82222,
                    "Minimum Price Alfa": 82222,
                    "Maximum Price Alfa": 82222,
                    "Minimum Price Digi": 82222,
                    "Maximum Price Digi": 82222,
                    "Final Price Alfa": 82999,
                    "Final Price Digi": 82999,
                    "Alfa Real Profit": 1199,
                    "Digi Real Profit": 1199,
                    "Alfa Reason": "MONOPOLY_PRICING",
                    "Digi Reason": "MONOPOLY_PRICING",
                    "Alfa Risk Alert": "NONE",
                    "Digi Risk Alert": "NONE",
                },
            ]
        )

        manifest = export_results(
            df=df,
            source_filename="test_catalog.csv",
            output_dir=tmp_path,
            timestamp=datetime(2026, 10, 26, 19, 0, 0),
            render_panel=False,
        )

        # Open generated workbook
        assert manifest.timestamped_xlsx_path is not None
        wb = openpyxl.load_workbook(manifest.timestamped_xlsx_path)
        ws = wb.active
        assert ws is not None

        col_names = [cell.value for cell in ws[1]]
        ship_idx = col_names.index("Shipping Estimate") + 1

        # Row 2 (Product with dimensions): has dynamic formula starting with =ROUNDUP
        ship_cell_with_dims = ws.cell(row=2, column=ship_idx)
        assert str(ship_cell_with_dims.value).startswith("=ROUNDUP")

        # Row 3 (Product with missing dimensions): retains fallback 3500, NOT overwritten with =ROUNDUP(0)
        ship_cell_fallback = ws.cell(row=3, column=ship_idx)
        assert ship_cell_fallback.value == 3500.0 or ship_cell_fallback.value == 3500

